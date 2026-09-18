import os
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def create_styled_document():
    doc = Document()
    
    # Page setup
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    PRIMARY = RGBColor(0, 74, 198)     # #004ac6
    SECONDARY = RGBColor(30, 41, 59)   # #1e293b
    MUTED = RGBColor(100, 116, 139)    # #64748b

    # Base styling
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Segoe UI'
    normal_style.font.size = Pt(10)
    normal_style.font.color.rgb = SECONDARY

    def set_cell_background(cell, fill_hex):
        tcPr = cell._tc.get_or_add_tcPr()
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
        tcPr.append(shd)

    def add_callout(quote_text):
        table = doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.columns[0].width = Inches(6.8)
        cell = table.cell(0, 0)
        set_cell_background(cell, 'F1F5F9')
        
        tcPr = cell._tc.get_or_add_tcPr()
        borders = parse_xml(f'''
            <w:tcBorders {nsdecls("w")}>
                <w:top w:val="none"/>
                <w:left w:val="single" w:sz="24" w:space="0" w:color="004AC6"/>
                <w:bottom w:val="none"/>
                <w:right w:val="none"/>
            </w:tcBorders>
        ''')
        tcPr.append(borders)
        
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.left_indent = Inches(0.15)
        p.paragraph_format.right_indent = Inches(0.15)
        run = p.add_run(f'“{quote_text}”')
        run.italic = True
        run.font.size = Pt(9.5)
        run.font.color.rgb = RGBColor(15, 23, 42)
        doc.add_paragraph()

    def add_code_block(code_text):
        table = doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.columns[0].width = Inches(6.8)
        cell = table.cell(0, 0)
        set_cell_background(cell, 'F8FAFC')
        
        tcPr = cell._tc.get_or_add_tcPr()
        borders = parse_xml(f'''
            <w:tcBorders {nsdecls("w")}>
                <w:top w:val="single" w:sz="6" w:space="0" w:color="CBD5E1"/>
                <w:left w:val="single" w:sz="6" w:space="0" w:color="CBD5E1"/>
                <w:bottom w:val="single" w:sz="6" w:space="0" w:color="CBD5E1"/>
                <w:right w:val="single" w:sz="6" w:space="0" w:color="CBD5E1"/>
            </w:tcBorders>
        ''')
        tcPr.append(borders)
        
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.left_indent = Inches(0.15)
        p.paragraph_format.right_indent = Inches(0.15)
        run = p.add_run(code_text)
        run.font.name = 'Consolas'
        run.font.size = Pt(8.5)
        run.font.color.rgb = RGBColor(15, 23, 42)
        doc.add_paragraph()

    # Document Header / Banner
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(10)
    title_p.paragraph_format.space_after = Pt(2)
    run_brand = title_p.add_run("OneAIAssist Platform Specification\n")
    run_brand.font.size = Pt(11)
    run_brand.font.bold = True
    run_brand.font.color.rgb = PRIMARY

    run_title = title_p.add_run("Multi-Channel Expansion: User Requirements, Technical Flows & Prompts Specification")
    run_title.font.size = Pt(20)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(15, 23, 42)

    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_after = Pt(18)
    run_sub = sub_p.add_run("Complete technical guide covering WhatsApp, Apple iMessage / SMS, and Meta Instagram Direct channels, including real visual architecture diagrams, user requests log, and production AI prompts.")
    run_sub.font.size = Pt(10.5)
    run_sub.font.color.rgb = MUTED

    # Read markdown source
    md_path = 'docs/multi_channel_flows_and_prompts.md'
    with open(md_path, 'r', encoding='utf-8') as f:
        content = f.read()

    # Split into sections
    lines = content.split('\n')
    i = 0
    in_code = False
    code_lang = ''
    code_lines = []

    in_table = False
    table_lines = []

    def flush_table(t_lines):
        if not t_lines:
            return
        rows_data = []
        for line in t_lines:
            if re.match(r'^\s*\|?\s*:?-+:?\s*\|', line):
                continue  # header divider line
            parts = [p.strip() for p in line.strip().strip('|').split('|')]
            rows_data.append(parts)
        
        if not rows_data:
            return

        cols = len(rows_data[0])
        table = doc.add_table(rows=len(rows_data), cols=cols)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        
        for r_idx, r_data in enumerate(rows_data):
            row = table.rows[r_idx]
            is_header = (r_idx == 0)
            for c_idx, val in enumerate(r_data):
                if c_idx < cols:
                    cell = row.cells[c_idx]
                    set_cell_background(cell, '004AC6' if is_header else ('F8FAFC' if r_idx % 2 == 1 else 'FFFFFF'))
                    p = cell.paragraphs[0]
                    p.paragraph_format.space_before = Pt(4)
                    p.paragraph_format.space_after = Pt(4)
                    
                    # Clean markdown bold/italic
                    clean_text = val.replace('**', '').replace('*', '')
                    run = p.add_run(clean_text)
                    run.font.size = Pt(8.5 if not is_header else 9)
                    run.font.bold = is_header
                    run.font.color.rgb = RGBColor(255, 255, 255) if is_header else SECONDARY
        doc.add_paragraph()

    while i < len(lines):
        line = lines[i]

        # Check code fence
        if line.startswith('```'):
            if not in_code:
                in_code = True
                code_lang = line[3:].strip()
                code_lines = []
            else:
                in_code = False
                joined_code = '\n'.join(code_lines)
                
                # Check if this was a mermaid diagram
                if code_lang == 'mermaid':
                    if 'subgraph Inbound Channels' in joined_code:
                        p_diag = doc.add_paragraph()
                        p_diag.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        p_diag.paragraph_format.space_before = Pt(8)
                        doc.add_picture('docs/images/multi_channel_architecture.png', width=Inches(6.5))
                        p_cap = doc.add_paragraph()
                        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        r_cap = p_cap.add_run("Figure 1: Multi-Channel Ingress, Queue & Outbound Architecture")
                        r_cap.font.size = Pt(9)
                        r_cap.italic = True
                        r_cap.font.color.rgb = MUTED
                        doc.add_paragraph()
                    elif 'stateDiagram-v2' in joined_code or 'NEW_LEAD' in joined_code:
                        p_diag = doc.add_paragraph()
                        p_diag.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        p_diag.paragraph_format.space_before = Pt(8)
                        doc.add_picture('docs/images/state_transition_diagram.png', width=Inches(6.0))
                        p_cap = doc.add_paragraph()
                        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        r_cap = p_cap.add_run("Figure 2: Stateful Lead Qualification & Compliance State Machine")
                        r_cap.font.size = Pt(9)
                        r_cap.italic = True
                        r_cap.font.color.rgb = MUTED
                        doc.add_paragraph()
                    else:
                        add_code_block(joined_code)
                else:
                    add_code_block(joined_code)
            i += 1
            continue

        if in_code:
            code_lines.append(line)
            i += 1
            continue

        # Check tables
        if line.strip().startswith('|') and '|' in line.strip()[1:]:
            if not in_table:
                in_table = True
                table_lines = [line]
            else:
                table_lines.append(line)
            i += 1
            continue
        elif in_table:
            in_table = False
            flush_table(table_lines)
            table_lines = []

        # Check blockquote / user prompt
        if line.startswith('> '):
            quote_lines = [line[2:].strip()]
            while i + 1 < len(lines) and (lines[i + 1].startswith('> ') or lines[i + 1].startswith('>')):
                i += 1
                q_sub = lines[i][2:].strip() if lines[i].startswith('> ') else lines[i][1:].strip()
                quote_lines.append(q_sub)
            raw_quote = ' '.join(quote_lines).replace('*', '')
            add_callout(raw_quote)
            i += 1
            continue

        # Headings
        if line.startswith('# ') and not line.startswith('## '):
            i += 1
            continue
        elif line.startswith('## '):
            h_text = line[3:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(16)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(h_text)
            run.font.size = Pt(14)
            run.font.bold = True
            run.font.color.rgb = PRIMARY
        elif line.startswith('### '):
            h_text = line[4:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(h_text)
            run.font.size = Pt(12)
            run.font.bold = True
            run.font.color.rgb = SECONDARY
        elif line.startswith('#### '):
            h_text = line[5:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(h_text)
            run.font.size = Pt(10.5)
            run.font.bold = True
            run.font.color.rgb = SECONDARY
        elif line.strip() == '---':
            pass
        elif line.strip().startswith('- ') or line.strip().startswith('* '):
            bullet_text = line.strip()[2:]
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            
            # Simple bold handler inside bullets
            parts = re.split(r'(\*\*.*?\*\*)', bullet_text)
            for pt in parts:
                if pt.startswith('**') and pt.endswith('**'):
                    r = p.add_run(pt[2:-2])
                    r.bold = True
                else:
                    p.add_run(pt)
        elif re.match(r'^\d+\.\s+', line.strip()):
            match = re.match(r'^\d+\.\s+(.*)', line.strip())
            num_text = match.group(1) if match else line.strip()
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            parts = re.split(r'(\*\*.*?\*\*)', num_text)
            for pt in parts:
                if pt.startswith('**') and pt.endswith('**'):
                    r = p.add_run(pt[2:-2])
                    r.bold = True
                else:
                    p.add_run(pt)
        elif line.strip():
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(4)
            parts = re.split(r'(\*\*.*?\*\*)', line)
            for pt in parts:
                if pt.startswith('**') and pt.endswith('**'):
                    r = p.add_run(pt[2:-2])
                    r.bold = True
                else:
                    p.add_run(pt)

        i += 1

    if in_table:
        flush_table(table_lines)

    output_path = 'docs/Multi_Channel_Flows_and_Prompts.docx'
    doc.save(output_path)
    print(f"Successfully generated Word document: {output_path} (Size: {os.path.getsize(output_path)} bytes)")

if __name__ == '__main__':
    create_styled_document()
