import os
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)

def build_docx(md_path, docx_path, images_base_dir):
    doc = Document()
    
    # Configure page margins (0.75 inch)
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)
        
    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()
        
    in_code_block = False
    code_lines = []
    
    for line in lines:
        raw_line = line.rstrip('\n')
        
        # Handle code block demarcation
        if raw_line.startswith('```'):
            if in_code_block:
                # Flush code block
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Inches(0.2)
                run = p.add_run('\n'.join(code_lines))
                run.font.name = 'Consolas'
                run.font.size = Pt(9.5)
                run.font.color.rgb = RGBColor(0x24, 0x29, 0x2E)
                code_lines = []
                in_code_block = False
            else:
                in_code_block = True
            continue
            
        if in_code_block:
            code_lines.append(raw_line)
            continue
            
        # Heading 1
        if raw_line.startswith('# '):
            h = doc.add_heading(raw_line[2:], level=1)
            h.paragraph_format.space_before = Pt(12)
            h.paragraph_format.space_after = Pt(6)
            for r in h.runs:
                r.font.name = 'Segoe UI'
                r.font.color.rgb = RGBColor(0x1B, 0x4B, 0x91)
            continue
            
        # Heading 2
        if raw_line.startswith('## '):
            h = doc.add_heading(raw_line[3:], level=2)
            h.paragraph_format.space_before = Pt(10)
            h.paragraph_format.space_after = Pt(4)
            for r in h.runs:
                r.font.name = 'Segoe UI'
                r.font.color.rgb = RGBColor(0x00, 0x4A, 0xC6)
            continue
            
        # Heading 3
        if raw_line.startswith('### '):
            h = doc.add_heading(raw_line[4:], level=3)
            h.paragraph_format.space_before = Pt(8)
            h.paragraph_format.space_after = Pt(3)
            for r in h.runs:
                r.font.name = 'Segoe UI'
                r.font.color.rgb = RGBColor(0x1C, 0x1B, 0x1F)
            continue

        # Horizontal Rule
        if raw_line == '---':
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run('_________________________________________________________________________________')
            run.font.color.rgb = RGBColor(0xCC, 0xD0, 0xD9)
            continue

        # Blockquote Alert
        if raw_line.startswith('> '):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.3)
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(3)
            text = raw_line[2:].strip()
            run = p.add_run(text)
            run.font.italic = True
            run.font.size = Pt(10)
            run.font.color.rgb = RGBColor(0x4A, 0x55, 0x68)
            continue

        # Image embed ![caption](./images/filename.png)
        img_match = re.search(r'!\[(.*?)\]\((.*?)\)', raw_line)
        if img_match:
            caption, img_rel_path = img_match.groups()
            img_rel_path = img_rel_path.replace('./images/', '').replace('file:///', '')
            img_abs_path = os.path.join(images_base_dir, os.path.basename(img_rel_path))
            
            if os.path.exists(img_abs_path):
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_before = Pt(6)
                p.paragraph_format.space_after = Pt(2)
                try:
                    p.add_run().add_picture(img_abs_path, width=Inches(6.0))
                except Exception as e:
                    p.add_run(f'[Image Error: {e}]')
                
                # Caption line
                if caption:
                    cp = doc.add_paragraph()
                    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    cp.paragraph_format.space_before = Pt(0)
                    cp.paragraph_format.space_after = Pt(8)
                    crun = cp.add_run(caption)
                    crun.font.italic = True
                    crun.font.size = Pt(9)
                    crun.font.color.rgb = RGBColor(0x71, 0x80, 0x96)
            continue

        # List items
        if raw_line.startswith('- ') or raw_line.startswith('* '):
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run(raw_line[2:])
            run.font.size = Pt(10.5)
            continue

        # Normal text paragraph
        if raw_line.strip():
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run(raw_line)
            run.font.size = Pt(10.5)

    doc.save(docx_path)
    print(f'Successfully generated Word document at: {docx_path}')

if __name__ == '__main__':
    images_dir = r'D:\codebase\oneaiassist_v1\docs\images'
    
    # 1. Analytics & Reports Guide Word Doc
    build_docx(
        r'D:\codebase\oneaiassist_v1\docs\analytics_and_reports_guide.md',
        r'D:\codebase\oneaiassist_v1\docs\Analytics_and_Reports_Guide.docx',
        images_dir
    )
    
    # 2. Updated Admin User Guide Word Doc
    build_docx(
        r'D:\codebase\oneaiassist_v1\docs\admin_user_guide.md',
        r'D:\codebase\oneaiassist_v1\docs\ADMIN_USER_GUIDE_V2.docx',
        images_dir
    )
