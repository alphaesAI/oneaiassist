import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def main():
    doc = docx.Document()

    # Set Margins (1 inch all around)
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    COLOR_PRIMARY = RGBColor(0, 74, 198)
    COLOR_SECONDARY = RGBColor(28, 27, 31)
    COLOR_MUTED = RGBColor(73, 69, 79)

    # Title
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_p.add_run("OneAI Assist — Admin User Guide & Operational Manual")
    title_run.font.name = "Arial"
    title_run.font.size = Pt(22)
    title_run.font.bold = True
    title_run.font.color.rgb = COLOR_PRIMARY

    sub_p = doc.add_paragraph()
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_run = sub_p.add_run("Complete Administrator Documentation, Credentials & Verification Screenshots")
    sub_run.font.name = "Arial"
    sub_run.font.size = Pt(11)
    sub_run.font.italic = True
    sub_run.font.color.rgb = COLOR_MUTED

    doc.add_paragraph() # spacing

    # --- Credentials Callout Table ---
    cred_table = doc.add_table(rows=5, cols=2)
    cred_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cred_table.autofit = False

    row_data = [
        ("Login URL", "http://localhost:3000/login"),
        ("Admin Email", "admin@primemarketingexperts.com"),
        ("Default Password", "password123"),
        ("Assigned Role", "ADMIN (Agency Administrator)"),
        ("Tenant Workspace", "Prime Marketing Experts (tenant_pme_ff9xl)")
    ]

    for i, (k, v) in enumerate(row_data):
        row_cells = cred_table.rows[i].cells
        row_cells[0].text = k
        row_cells[1].text = v
        
        p0 = row_cells[0].paragraphs[0]
        p0.runs[0].font.bold = True
        p0.runs[0].font.name = "Arial"
        p0.runs[0].font.size = Pt(10)
        p0.runs[0].font.color.rgb = COLOR_SECONDARY
        
        p1 = row_cells[1].paragraphs[0]
        p1.runs[0].font.name = "Arial"
        p1.runs[0].font.size = Pt(10)
        if i in [0, 1, 2]:
            p1.runs[0].font.bold = True
            p1.runs[0].font.color.rgb = COLOR_PRIMARY

        set_cell_background(row_cells[0], "F4F6FA")
        set_cell_background(row_cells[1], "FFFFFF")
        set_cell_margins(row_cells[0], 120, 120, 150, 150)
        set_cell_margins(row_cells[1], 120, 120, 150, 150)

    doc.add_paragraph()

    # --- Helper to add sections ---
    img_dir = r"D:\codebase\oneaiassist_v1\docs\images"

    def add_section(number, title, text_paragraphs, image_filename=None, image_caption=None):
        h = doc.add_heading(level=1)
        h_run = h.add_run(f"{number}. {title}")
        h_run.font.name = "Arial"
        h_run.font.size = Pt(15)
        h_run.font.bold = True
        h_run.font.color.rgb = COLOR_PRIMARY

        for p_text in text_paragraphs:
            p = doc.add_paragraph()
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.space_after = Pt(6)
            p_run = p.add_run(p_text)
            p_run.font.name = "Arial"
            p_run.font.size = Pt(10.5)
            p_run.font.color.rgb = COLOR_SECONDARY

        if image_filename:
            img_path = os.path.join(img_dir, image_filename)
            if os.path.exists(img_path):
                img_p = doc.add_paragraph()
                img_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                img_p.paragraph_format.space_before = Pt(8)
                img_p.paragraph_format.space_after = Pt(4)
                
                img_p.add_run().add_picture(img_path, width=Inches(5.8))
                
                if image_caption:
                    cap_p = doc.add_paragraph()
                    cap_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    cap_p.paragraph_format.space_after = Pt(12)
                    cap_run = cap_p.add_run(f"Figure {number}: {image_caption}")
                    cap_run.font.name = "Arial"
                    cap_run.font.size = Pt(9)
                    cap_run.font.italic = True
                    cap_run.font.color.rgb = COLOR_MUTED
            doc.add_paragraph()

    # 1. System Authentication
    add_section(
        1,
        "System Access & Credentials",
        [
            "To access the OneAI Assist workspace as an agency administrator, navigate to http://localhost:3000/login in any web browser.",
            "Log in using email: admin@primemarketingexperts.com and password: password123. Administrators have full access to workspace settings, AI bot tuning, WhatsApp templates, and lead management."
        ],
        "login_page.png",
        "Login Screen with OneAIAssist Logo and Credentials Form"
    )

    # 2. Dashboard Overview
    add_section(
        2,
        "Dashboard Overview & Key Performance Metrics",
        [
            "The main agency overview page at http://localhost:3000/dashboard provides real-time visibility into active customer conversations, monthly leads, policies sold, and bot resolution metrics.",
            "All KPI cards and activity feed items are dynamically loaded from underlying PostgreSQL database records for Prime Marketing Experts."
        ],
        "dashboard_page.png",
        "Agency Dashboard Overview displaying live KPI metrics and branding logo"
    )

    # 3. WhatsApp Integration
    add_section(
        3,
        "WhatsApp Business Integration & Engine Provider Architecture",
        [
            "The Agency Settings page at http://localhost:3000/dashboard/settings enables administrators to connect their WhatsApp Business phone number using QR code scanning or 8-character pairing codes.",
            "Engine Architecture: OneAIAssist natively integrates two Node.js WhatsApp engines — OpenWA (@open-wa/wa-automate Chromium engine) and Baileys (@whiskeysockets/baileys WebSocket engine). (Note: Third-party tools like openclaw are not used).",
            "Once connected, the engine handles automated lead intake, RAG policy Q&A, and live agent handoffs seamlessly."
        ],
        "settings_page.png",
        "Agency Settings — WhatsApp Business Connection with Scannable QR Code"
    )

    # 4. Bot Studio & Flow Builder
    add_section(
        4,
        "AI Bot Studio & Flow Builder",
        [
            "The Bot Configuration Studio at http://localhost:3000/dashboard/bot-config features a drag-and-drop intake question flowchart canvas.",
            "Admins can configure validation rules, required fields, and conditional skip logic branching for automated customer qualification."
        ],
        "bot_config_canvas.png",
        "Bot Configuration Studio — Interactive Intake Question Flow Builder Canvas"
    )

    # 5. Knowledge Base & RAG Indexing
    add_section(
        5,
        "Knowledge Base & Vector RAG Document Index",
        [
            "The Knowledge Base tab under http://localhost:3000/dashboard/bot-config allows administrators to upload insurance policy PDFs (e.g. Apex Health Care plans).",
            "PDF text is extracted, chunked, and embedded into pgvector / Pinecone vectorstore using Google text-embedding-004 to power intelligent customer answers."
        ],
        "knowledge_base_tab.png",
        "Knowledge Base Tab displaying indexed U.S. Health Insurance policy documents"
    )

    # 6. Product Catalog
    add_section(
        6,
        "Insurance Product Catalog",
        [
            "The Product Catalog tab lists all managed health insurance products, states covered, monthly premium ranges, and maximum sum insured amounts.",
            "Data is dynamically fetched from database PolicyCatalogItem records to ensure accurate quote generation."
        ],
        "product_catalog_tab.png",
        "Insurance Product Catalog displaying managed policies and monthly premiums"
    )

    # 7. WhatsApp Template Manager
    add_section(
        7,
        "WhatsApp Template Manager",
        [
            "The Template Manager at http://localhost:3000/dashboard/templates allows creating and approving marketing, utility, and authentication message templates with variable placeholders."
        ],
        "templates_page.png",
        "WhatsApp Template Manager displaying approved message templates"
    )

    # 8. Broadcast Campaign Center
    add_section(
        8,
        "Broadcast Campaign Center",
        [
            "The Broadcast Center at http://localhost:3000/dashboard/broadcast lets admins launch targeted bulk WhatsApp messaging campaigns based on customer tags and pipeline stages."
        ],
        "broadcast_center.png",
        "Broadcast Campaign Center for target bulk messaging"
    )

    # 9. Lead Pipeline Management
    add_section(
        9,
        "Lead Pipeline & CRM Management",
        [
            "The Lead Pipeline at http://localhost:3000/dashboard/leads tracks customer opportunities from initial inquiry to policy activation in dual Kanban and table views."
        ],
        "leads_add_modal.png",
        "Lead Pipeline Kanban Board with + Add Lead Dialog"
    )

    # 10. Customer Directory & 360° Profile
    add_section(
        10,
        "Customer Directory & 360° Profile View",
        [
            "The Customer Directory at http://localhost:3000/dashboard/customers maintains all customer contact records, opt-in consent statuses, and historical policy engagements."
        ],
        "customer_profile.png",
        "Customer 360° Profile View with policy history"
    )

    out_dir = r"D:\codebase\oneaiassist_v1\docs"
    out_path = os.path.join(out_dir, "ADMIN_USER_GUIDE.docx")
    try:
        doc.save(out_path)
        print(f"Successfully generated Word document at: {out_path}")
    except PermissionError:
        alt_path = os.path.join(out_dir, "ADMIN_USER_GUIDE_V2.docx")
        doc.save(alt_path)
        print(f"Primary file locked. Successfully generated Word document at: {alt_path}")

if __name__ == "__main__":
    main()
